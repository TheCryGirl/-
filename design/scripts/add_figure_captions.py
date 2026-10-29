# -*- coding: utf-8 -*-
"""给 docx 的图片补图题（图片下方一行「图 X-Y 标题」）。

源稿里的图题写在 Markdown 的 ![图 2-1 系统逻辑架构](path) 里，
但 office_write_docx 生成时把 alt 文本丢了，导致 docx 里没有图题。
本脚本按图题顺序，依次补到对应图片段落之后，并做居中、加粗、无缩进。

用法：python add_figure_captions.py <docx路径> <图题1> <图题2> ...
"""
import sys
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from copy import deepcopy

PATH = sys.argv[1]
CAPTIONS = sys.argv[2:]

doc = Document(PATH)
body = doc.element.body

# 找出所有含图片的段落
img_paras = []
for p in doc.paragraphs:
    if 'graphic' in p._p.xml:
        img_paras.append(p)

if len(img_paras) != len(CAPTIONS):
    print(f'警告: 图片 {len(img_paras)} 个，图题 {len(CAPTIONS)} 条，数量不匹配')
    sys.exit(1)

added = 0
for img_p, caption in zip(img_paras, CAPTIONS):
    # 已经有图题就跳过
    nxt = img_p._p.getnext()
    if nxt is not None and nxt.tag.endswith('}p'):
        from docx.text.paragraph import Paragraph
        if Paragraph(nxt, img_p._parent).text.strip().startswith('图'):
            continue

    # 用图片段落复制一个空段，插到它后面，保证样式继承
    new_p = deepcopy(img_p._p)
    # 清掉复制来的图片
    for child in list(new_p):
        if child.tag.endswith('}r'):
            new_p.remove(child)
    img_p._p.addnext(new_p)

    from docx.text.paragraph import Paragraph
    para = Paragraph(new_p, img_p._parent)
    run = para.add_run(caption)
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(12)
    run.font.bold = True

    pf = para.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(3)
    pf.space_after = Pt(6)
    added += 1

doc.save(PATH)
print(f'已补图题 {added} 条')
