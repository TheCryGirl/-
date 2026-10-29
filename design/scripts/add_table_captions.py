# -*- coding: utf-8 -*-
"""给 docx 的表格补表题（表格上方一行「表 X-Y 标题」）。

指导书要求表题在表格上方、居中。实验一的表题是「表 1-1 目标用户与核心价值」这种格式。

用法：python add_table_captions.py <docx路径> <表题1> <表题2> ...
      表题按文档中表格出现的顺序对应。
"""
import sys
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from copy import deepcopy
from docx.text.paragraph import Paragraph
from docx.table import Table

PATH = sys.argv[1]
CAPTIONS = sys.argv[2:]

doc = Document(PATH)

tables = []
for child in doc.element.body.iterchildren():
    if child.tag.endswith('}tbl'):
        tables.append(child)

if len(tables) != len(CAPTIONS):
    print(f'警告: 表格 {len(tables)} 个，表题 {len(CAPTIONS)} 条，数量不匹配')
    sys.exit(1)

added = 0
for tbl, caption in zip(tables, CAPTIONS):
    # 表格上方已有表题则跳过
    prev = tbl.getprevious()
    if prev is not None and prev.tag.endswith('}p'):
        if Paragraph(prev, doc).text.strip().startswith('表'):
            continue

    # 复制表格后面的第一段做样式模板；没有就新建
    new_p = None
    nxt = tbl.getnext()
    if nxt is not None and nxt.tag.endswith('}p'):
        new_p = deepcopy(nxt)
        for child in list(new_p):
            if child.tag.endswith('}r') or child.tag.endswith('}hyperlink'):
                new_p.remove(child)
    else:
        from docx.oxml import OxmlElement
        new_p = OxmlElement('w:p')

    tbl.addprevious(new_p)
    para = Paragraph(new_p, doc)
    # 清空可能残留的文字
    for r in list(para.runs):
        r._element.getparent().remove(r._element)

    run = para.add_run(caption)
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(12)
    run.font.bold = True

    pf = para.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    added += 1

doc.save(PATH)
print(f'已补表题 {added} 条')
