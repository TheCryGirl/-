# -*- coding: utf-8 -*-
"""把实验二报告正文排版对齐实验一的标准：
- 正文：宋体/Times New Roman 12pt，1.5 倍行距，首行缩进 24pt，两端对齐
- 标题：黑体，各级字号（H1 16 / H2 14 / H3 12 / H4 12），1.5 倍行距，不缩进
- 列表段落：同正文，但保留项目符号
- 表格首行加粗表头

只处理正文段落，不动封面与图片。
"""
import sys
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

PATH = sys.argv[1]
doc = Document(PATH)


def set_font(run, cn="宋体", ascii_="Times New Roman", size=12, bold=None):
    run.font.name = ascii_
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


HEAD_SIZES = {"Heading 1": 16, "Heading 2": 14, "Heading 3": 12, "Heading 4": 12}

n_body = n_head = 0
started = False
for p in doc.paragraphs:
    style = p.style.name
    text = p.text.strip()
    if not text:
        continue

    # 跳过封面与目录：从第一个 Heading 标题才开始排版
    if not started:
        if style.startswith("Heading"):
            started = True
        else:
            continue

    if style in HEAD_SIZES:
        size = HEAD_SIZES[style]
        for r in p.runs:
            set_font(r, cn="黑体", ascii_="Times New Roman", size=size, bold=False)
        pf = p.paragraph_format
        pf.space_before = Pt(12 if style == "Heading 1" else 6)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.5
        pf.first_line_indent = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        n_head += 1
        continue

    # 标题式段落（"图 2-1 ..." / "表 2-1 ..."）居中加粗
    if text.startswith("图 ") or text.startswith("表 "):
        for r in p.runs:
            set_font(r, cn="宋体", ascii_="Times New Roman", size=12, bold=True)
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Pt(0)
        pf.line_spacing = 1.15
        continue

    # 正文 / 列表
    for r in p.runs:
        set_font(r, cn="宋体", ascii_="Times New Roman", size=12)
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Pt(24)
    n_body += 1

# 表格内字号统一五号（10.5pt）、表头加粗
for t in doc.tables:
    for i, row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_font(r, cn="宋体", ascii_="Times New Roman",
                             size=10.5, bold=(i == 0))

doc.save(PATH)
print(f"body={n_body} headings={n_head} tables={len(doc.tables)}")
